"""Lógica de processamento dos slides, sem nenhuma dependência de interface.

Pode ser usada pela janela (app/ui.py), por linha de comando ou por
outros scripts. Cada etapa é uma função pequena para facilitar incluir
novas operações (ex.: trocar logo, rodapé, etc.) ao lado de `apply_background`.
"""

import os
from dataclasses import dataclass, field
from typing import Callable, Optional

from pptx import Presentation
from pptx.exc import PackageNotFoundError

# Nome dado à imagem inserida. Permite reconhecer e substituir um fundo
# aplicado anteriormente, em vez de empilhar imagens a cada execução.
BACKGROUND_SHAPE_NAME = "AutomacaoSlides_Fundo"

IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff")
PPTX_EXTENSION = ".pptx"


@dataclass
class Job:
    source_folder: str
    image_path: str
    output_folder: str
    include_subfolders: bool = False
    replace_previous_background: bool = True


@dataclass
class Result:
    total: int = 0
    processed: list = field(default_factory=list)
    failed: list = field(default_factory=list)  # (caminho, mensagem)
    cancelled: bool = False


ProgressCallback = Callable[[int, int, str, Optional[str]], None]


def _same_path(a: str, b: str) -> bool:
    return os.path.normcase(os.path.abspath(a)) == os.path.normcase(os.path.abspath(b))


def _is_inside(path: str, parent: str) -> bool:
    path = os.path.normcase(os.path.abspath(path))
    parent = os.path.normcase(os.path.abspath(parent))
    return path == parent or path.startswith(parent + os.sep)


def find_presentations(folder: str, include_subfolders: bool = False,
                       exclude_folder: Optional[str] = None) -> list:
    """Lista os .pptx da pasta (ignora arquivos temporários '~$' do PowerPoint)."""
    if not folder or not os.path.isdir(folder):
        return []

    def wanted(name: str) -> bool:
        return name.lower().endswith(PPTX_EXTENSION) and not name.startswith("~$")

    if not include_subfolders:
        return sorted(
            os.path.join(folder, n) for n in os.listdir(folder)
            if wanted(n) and os.path.isfile(os.path.join(folder, n))
        )

    found = []
    for root, dirs, files in os.walk(folder):
        if exclude_folder:
            dirs[:] = [d for d in dirs if not _is_inside(os.path.join(root, d), exclude_folder)]
        dirs.sort()
        found.extend(os.path.join(root, n) for n in sorted(files) if wanted(n))
    return found


def validate(job: Job) -> list:
    """Retorna a lista de problemas encontrados (vazia = pode processar)."""
    errors = []
    if not job.source_folder:
        errors.append("Selecione a pasta dos slides.")
    elif not os.path.isdir(job.source_folder):
        errors.append("A pasta dos slides não existe ou não está acessível.")

    if not job.image_path:
        errors.append("Selecione a imagem de fundo.")
    elif not os.path.isfile(job.image_path):
        errors.append("A imagem de fundo não foi encontrada.")
    elif not job.image_path.lower().endswith(IMAGE_EXTENSIONS):
        errors.append("Formato de imagem não suportado.")

    if not job.output_folder:
        errors.append("Selecione a pasta de saída.")
    elif job.source_folder and _same_path(job.source_folder, job.output_folder):
        errors.append("A pasta de saída deve ser diferente da pasta dos slides "
                      "(para não sobrescrever os originais).")
    return errors


def apply_background(prs, image_path: str, replace_previous: bool = True) -> None:
    """Insere a imagem cobrindo cada slide, atrás de todos os outros elementos."""
    for slide in prs.slides:
        tree = slide.shapes._spTree
        if replace_previous:
            for shape in list(slide.shapes):
                if shape.name == BACKGROUND_SHAPE_NAME:
                    tree.remove(shape._element)

        picture = slide.shapes.add_picture(
            image_path, 0, 0, width=prs.slide_width, height=prs.slide_height
        )
        picture._element.nvPicPr.cNvPr.name = BACKGROUND_SHAPE_NAME
        tree.remove(picture._element)
        tree.insert(2, picture._element)  # 0 e 1 são propriedades do grupo


def process_file(src: str, dst: str, image_path: str, replace_previous: bool = True) -> None:
    prs = Presentation(src)
    apply_background(prs, image_path, replace_previous)
    os.makedirs(os.path.dirname(dst) or ".", exist_ok=True)
    prs.save(dst)


def run(job: Job, on_progress: Optional[ProgressCallback] = None,
        should_cancel: Optional[Callable[[], bool]] = None) -> Result:
    """Processa todos os arquivos. Mantém a estrutura de subpastas na saída."""
    files = find_presentations(job.source_folder, job.include_subfolders,
                               exclude_folder=job.output_folder)
    result = Result(total=len(files))

    for index, src in enumerate(files, start=1):
        if should_cancel and should_cancel():
            result.cancelled = True
            break
        rel = os.path.relpath(src, job.source_folder)
        dst = os.path.join(job.output_folder, rel)
        error = None
        try:
            process_file(src, dst, job.image_path, job.replace_previous_background)
            result.processed.append(src)
        except PackageNotFoundError:
            error = "arquivo inválido ou corrompido"
            result.failed.append((src, error))
        except PermissionError:
            error = "sem permissão (o arquivo de saída pode estar aberto no PowerPoint)"
            result.failed.append((src, error))
        except Exception as exc:  # um arquivo com problema não interrompe o lote
            error = str(exc) or exc.__class__.__name__
            result.failed.append((src, error))
        if on_progress:
            on_progress(index, len(files), rel, error)
    return result
